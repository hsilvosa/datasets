"""Lossless-source extraction with explicit evidence and RAG offsets."""
from __future__ import annotations

import io
import re
import xml.etree.ElementTree as ET
from urllib.parse import urljoin

from bs4 import BeautifulSoup

from .archive_store import ArchiveConfig, Store, digest
from .io_utils import atomic_bytes
from .parser import xml_to_dict

PARSER_VERSION = "archive-2"
OCR_ENGINE = None
DOCUMENT_EXTENSIONS = (".pdf", ".xml", ".epub", ".csv", ".xls", ".xlsx", ".zip",
                       ".png", ".jpg", ".jpeg", ".gif", ".tif", ".tiff", ".svg",
                       ".txt", ".tsv", ".doc", ".docx", ".odt")


def text_of(node: ET.Element) -> str:
    if node.tag.split("}")[-1] == "table":
        rows = ["\t".join(" ".join(c.itertext()).strip() for c in row
                          if c.tag.split("}")[-1] in ("td", "th"))
                for row in node.iter() if row.tag.split("}")[-1] == "tr"]
        return "\n".join(rows)
    return " ".join("".join(node.itertext()).split())


def chunks(text: str, size: int, overlap: int):
    start = 0
    while start < len(text):
        end = min(start + size, len(text))
        yield start, end, text[start:end]
        if end == len(text):
            break
        start = end - overlap


def resource_kind(url: str) -> str:
    path = url.split("?", 1)[0].lower()
    if path.endswith(".pdf"):
        return "pdf"
    if path.endswith((".png", ".jpg", ".jpeg", ".gif", ".tif", ".tiff")):
        return "image"
    if path.endswith(".svg"):
        return "svg"
    if path.endswith(".xml") or "/xml.php" in path:
        return "xml"
    if path.endswith(".epub"):
        return "epub"
    if path.endswith((".html", ".htm", ".php")):
        return "html"
    return "asset"


class Extractor:
    def __init__(self, store: Store, config: ArchiveConfig, document_id: str,
                 url: str, checksum: str):
        self.store = store
        self.config = config
        self.document_id = document_id
        self.url = url
        self.checksum = checksum
        self.namespace = ""

    def emit(self, table: str, item: str, payload: dict) -> None:
        self.store.record(table, self.document_id, self.url, self.checksum,
                          PARSER_VERSION, self.namespace + item, payload)

    def full_text(self, text: str, method: str, item: str = "document", **metadata) -> None:
        self.emit("texts", item, {"text": text, "char_count": len(text),
                                  "method": method, **metadata})

    def rag(self, text: str, item: str, **metadata) -> None:
        for start, end, fragment in chunks(text, self.config.chunk_chars,
                                          self.config.chunk_overlap):
            self.emit("rag_chunks", f"{item}:{start}", {
                "text": fragment, "char_start": start, "char_end": end,
                "parent_id": item, "citation_url": self.url, **metadata,
            })

    def linked_asset(self, value: str, role: str = "attachment") -> None:
        url = urljoin(self.url, value)
        self.store.enqueue(url, resource_kind(url), 45, self.document_id, role)

    def xml(self, content: bytes) -> None:
        root = ET.fromstring(content)
        self.emit("metadata", "root_attributes", {"attributes": root.attrib})
        for name in ("metadatos", "metadata", "analisis"):
            node = root.find(f".//{name}")
            if node is not None:
                self.emit("metadata", name, {"section": name, "value": xml_to_dict(node),
                                            "xml": ET.tostring(node, encoding="unicode")})
        # All source metadata is archived, including fields not in a normalized schema.
        for node in root.iter():
            tag = node.tag.split("}")[-1]
            if tag.startswith("url_") and node.text and node.text.strip():
                if tag == "url_eli":
                    self.emit("references", "eli", {"target": node.text.strip(), "role": "eli"})
                else:
                    self.linked_asset(node.text.strip(), tag)
            for attr in ("src", "href"):
                value = node.get(attr)
                if value:
                    if tag in ("img", "imagen", "image") or value.lower().split("?", 1)[0].endswith(DOCUMENT_EXTENSIONS):
                        self.linked_asset(value, "inline_image" if tag in ("img", "imagen", "image") else "attachment")
                    else:
                        self.emit("references", digest(value), {"target": urljoin(self.url, value),
                                                                 "label": text_of(node)})
        bodies = root.findall("./texto") + root.findall("./data/texto")
        if root.tag == "texto":
            bodies = [root]
        if not bodies:
            raise ValueError("Document XML has no texto element")
        for body_index, body in enumerate(bodies):
            complete = "\n".join(text_of(node) for node in body if text_of(node))
            self.full_text(complete, "xml", f"body:{body_index}",
                           contains_historical_versions=bool(body.findall("bloque/version")))
            if not complete.strip():
                self.store.issue(self.document_id, self.url, "empty_text", "XML body has no text")
            blocks = body.findall("bloque")
            if blocks:
                for index, block in enumerate(blocks):
                    versions = block.findall("version")
                    for version_index, version in enumerate(versions):
                        block_id = block.get("id") or str(index)
                        item = f"{body_index}:{block_id}:{version_index}"
                        text = "\n".join(text_of(child) for child in version)
                        metadata = {"block_id": block_id, "block_attributes": block.attrib,
                                    "version_attributes": version.attrib,
                                    "is_latest_version": version_index == len(versions) - 1,
                                    "version_index": version_index,
                                    "valid_from": version.get("fecha_vigencia"),
                                    "publication_date": version.get("fecha_publicacion")}
                        self.emit("blocks", item, {"text": text, "xml": ET.tostring(version, encoding="unicode"), **metadata})
                        self.rag(text, item, **metadata)
                        self.structure(version, item)
            else:
                self.structure(body, f"body:{body_index}")
                self.rag(complete, f"body:{body_index}", is_latest_version=True)
                if self.document_id.startswith("BORME-A-"):
                    self.borme(body)

    def structure(self, body: ET.Element, prefix: str) -> None:
        heading = None
        for index, child in enumerate(body):
            tag = child.tag.split("}")[-1]
            label = child.get("class", "")
            text = text_of(child)
            if any(word in label for word in ("articulo", "anexo", "capitulo", "titulo", "seccion")):
                heading = text
            self.emit("structure", f"{prefix}:{index}", {
                "order": index, "parent_id": prefix, "tag": tag,
                "attributes": child.attrib, "heading": heading, "text": text,
                "xml": ET.tostring(child, encoding="unicode"),
            })
        for index, table in enumerate(body.iter("table")):
            self.emit("tables", f"{prefix}:{index}", {
                "parent_id": prefix, "xml": ET.tostring(table, encoding="unicode"),
                "rows": [[" ".join(c.itertext()).strip() for c in row
                          if c.tag in ("td", "th")] for row in table.iter("tr")],
            })

    def borme(self, body: ET.Element) -> None:
        current: dict | None = None
        index = 0
        for child in body:
            text = text_of(child)
            match = re.match(r"^(\d+)\s*[-–]\s*(.+)$", text)
            if match and "articulo" in child.get("class", ""):
                if current:
                    self.emit_act(str(index), current)
                    index += 1
                current = {"announcement_number": match[1], "company": match[2],
                           "heading": text, "paragraphs": []}
            elif current:
                current["paragraphs"].append(text)
        if current:
            self.emit_act(str(index), current)
        else:
            self.store.issue(self.document_id, self.url, "borme_segmentation",
                             "No individual registry acts recognized; full text retained")

    def emit_act(self, item: str, act: dict) -> None:
        text = "\n".join([act["heading"], *act["paragraphs"]])
        evidence = []
        patterns = {
            "tax_id": r"\b[ABCDEFGHJNPQRSUVW]\d{7}[0-9A-J]\b",
            "registry": r"Datos registrales\.[^\n]+",
            "capital": r"Capital\s*:\s*[^.\n]*(?:\.[^\n]*Euros)?",
            "address": r"Domicilio\s*:\s*[^\n]+",
            "appointment": r"(?:Nombramientos|Ceses|Revocaciones|Reelecciones)\.[^\n]+",
            "corporate_purpose": r"Objeto social\s*:\s*[^\n]+",
        }
        for field, pattern in patterns.items():
            for match in re.finditer(pattern, text, re.IGNORECASE):
                evidence.append({"field": field, "raw_value": match[0],
                                 "start": match.start(), "end": match.end(),
                                 "method": "regex", "requires_validation": True})
        self.emit("borme_acts", item, {**act, "text": text, "fields": evidence})
        self.rag(text, f"act:{item}", company=act["company"],
                 announcement_number=act["announcement_number"], is_latest_version=True)

    def html(self, content: bytes) -> None:
        soup = BeautifulSoup(content, "html.parser")
        body = soup.select_one("#textoxslt, #textoxs, #texto, #textoboe, .documento, .texto")
        if body is None:
            # The HTML representation is not silently accepted as a document body.
            self.store.issue(self.document_id, self.url, "html_body",
                             "Document body selector not recognized; original retained")
            return
        for tag in body.select("img[src], a[href], source[src], object[data]"):
            value = tag.get("src") or tag.get("href") or tag.get("data")
            if not value:
                continue
            if tag.name != "a" or value.lower().split("?", 1)[0].endswith(DOCUMENT_EXTENSIONS):
                self.linked_asset(value, "inline_image" if tag.name == "img" else "attachment")
            else:
                self.emit("references", digest(value), {"target": urljoin(self.url, value),
                                                         "label": tag.get_text(" ", strip=True)})
        text = body.get_text("\n", strip=True)
        self.full_text(text, "html")
        self.rag(text, "html", is_latest_version=True)
        for index, table in enumerate(body.select("table")):
            self.emit("tables", f"html:{index}", {
                "html": str(table), "rows": [[cell.get_text(" ", strip=True)
                for cell in row.find_all(["td", "th"], recursive=False)]
                for row in table.select("tr")],
            })

    def ocr(self, image_bytes: bytes) -> tuple[str, list, float | None]:
        global OCR_ENGINE
        if OCR_ENGINE is None:
            from rapidocr_onnxruntime import RapidOCR
            OCR_ENGINE = RapidOCR(det_use_cuda=False, rec_use_cuda=False, cls_use_cuda=False)
        lines, _ = OCR_ENGINE(image_bytes)
        values = [{"bbox": line[0], "text": line[1], "confidence": float(line[2])}
                  for line in (lines or [])]
        mean = sum(v["confidence"] for v in values) / len(values) if values else None
        return "\n".join(v["text"] for v in values), values, mean

    def pdf(self, content: bytes) -> None:
        import pymupdf
        complete = []
        with pymupdf.open(stream=content, filetype="pdf") as pdf:
            self.emit("metadata", "pdf", {"metadata": pdf.metadata,
                                          "page_count": len(pdf), "toc": pdf.get_toc()})
            if pdf.needs_pass or not len(pdf):
                raise ValueError("Encrypted or empty PDF")
            for index, page in enumerate(pdf):
                text = page.get_text("text", sort=True)
                blocks = page.get_text("blocks", sort=True)
                images = page.get_images(full=True)
                ocr = []
                method = "pdf_native"
                # Scans can have headers but no searchable body: inspect image area too.
                image_area = sum(r.get_area() for entry in images
                                 for r in page.get_image_rects(entry[0]))
                needs_ocr = len(text.strip()) < 80 or image_area > page.rect.get_area() * 0.35
                if needs_ocr:
                    pixels = page.get_pixmap(dpi=self.config.ocr_dpi)
                    png = pixels.tobytes("png")
                    image_hash = digest(png)
                    path = self.store.root / "objects" / image_hash[:2] / image_hash
                    if not path.exists():
                        atomic_bytes(path, png)
                    ocr_text, ocr, confidence = self.ocr(png)
                    method = "pdf_native_and_ocr"
                    text = text + "\n" + ocr_text
                    self.store.issue(self.document_id, self.url, "ocr_review",
                                     f"Page {index + 1}: OCR requires validation (mean={confidence})")
                    self.emit("page_images", str(index), {
                        "page": index + 1, "sha256": image_hash,
                        "path": str(path.relative_to(self.store.root)),
                    })
                payload = {"page": index + 1, "text": text, "method": method,
                           "native_blocks": blocks, "ocr_lines": ocr,
                           "image_count": len(images), "width": page.rect.width,
                           "height": page.rect.height}
                self.emit("pages", str(index), payload)
                self.rag(text, f"page:{index + 1}", page=index + 1, method=method,
                         requires_review=needs_ocr, is_latest_version=True)
                if not text.strip():
                    self.store.issue(self.document_id, self.url, "empty_page", str(index + 1))
                complete.append(text)
                # Preserve embedded images as addressable assets, including non-text diagrams.
                for image_index, entry in enumerate(images):
                    image = pdf.extract_image(entry[0])
                    if not image:
                        continue
                    image_hash = digest(image["image"])
                    image_path = self.store.root / "objects" / image_hash[:2] / image_hash
                    if not image_path.exists():
                        atomic_bytes(image_path, image["image"])
                    self.emit("embedded_images", f"{index}:{image_index}", {
                        "page": index + 1, "sha256": image_hash,
                        "path": str(image_path.relative_to(self.store.root)),
                        "extension": image["ext"], "width": image["width"],
                        "height": image["height"],
                    })
                for table_index, table in enumerate(page.find_tables().tables):
                    self.emit("tables", f"pdf:{index}:{table_index}", {
                        "page": index + 1, "bbox": list(table.bbox),
                        "rows": table.extract(), "method": "pdf_geometry",
                        "requires_validation": True,
                    })
            for index in range(pdf.embfile_count()):
                attachment = pdf.embfile_get(index)
                checksum = digest(attachment)
                path = self.store.root / "objects" / checksum[:2] / checksum
                if not path.exists():
                    atomic_bytes(path, attachment)
                self.emit("attachments", f"embedded:{index}", {
                    "sha256": checksum, "path": str(path.relative_to(self.store.root)),
                    "metadata": pdf.embfile_info(index), "requires_validation": True,
                })
                self.store.issue(self.document_id, self.url, "embedded_attachment_review",
                                 f"Embedded attachment {index} retained; extraction needs review")
        self.full_text("\f".join(complete), "pdf", page_count=len(complete))

    def image(self, content: bytes) -> None:
        text, lines, confidence = self.ocr(content)
        self.full_text(text, "image_ocr", lines=lines, mean_confidence=confidence)
        self.rag(text, "image", method="image_ocr", requires_review=True)
        self.store.issue(self.document_id, self.url, "ocr_review", "Image OCR requires validation")

    def asset(self, content: bytes, kind: str) -> None:
        if kind == "svg":
            root = ET.fromstring(content)
            text = "\n".join("".join(n.itertext()) for n in root.iter()
                             if n.tag.split("}")[-1] == "text")
            self.full_text(text, "svg")
            self.rag(text, "svg", requires_review=True)
        elif kind == "epub":
            import zipfile
            with zipfile.ZipFile(io.BytesIO(content)) as book:
                for index, name in enumerate(book.namelist()):
                    if name.lower().endswith((".html", ".xhtml", ".htm")):
                        text = BeautifulSoup(book.read(name), "html.parser").get_text("\n", strip=True)
                        self.full_text(text, "epub", f"epub:{index}", member=name)
                        self.rag(text, f"epub:{index}", member=name)
        elif self.url.lower().split("?", 1)[0].endswith((".csv", ".tsv", ".txt")):
            from charset_normalizer import from_bytes
            decoded = from_bytes(content).best()
            if decoded is None:
                raise ValueError("Unable to identify attachment encoding")
            text = str(decoded)
            self.full_text(text, "text_attachment", encoding=decoded.encoding)
            self.rag(text, "text_attachment")
            if not self.url.lower().endswith(".txt"):
                import csv
                dialect = csv.Sniffer().sniff(text[:10000])
                for index, row in enumerate(csv.reader(io.StringIO(text), dialect)):
                    self.emit("tables", f"csv:{index}", {"row_index": index, "cells": row})
            return
        elif self.url.lower().split("?", 1)[0].endswith(".xlsx"):
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            try:
                for sheet in book:
                    texts = []
                    for index, row in enumerate(sheet.iter_rows(values_only=True)):
                        self.emit("tables", f"xlsx:{sheet.title}:{index}", {
                            "sheet": sheet.title, "row_index": index, "cells": list(row),
                        })
                        texts.append("\t".join(str(value) if value is not None else "" for value in row))
                    text = "\n".join(texts)
                    self.full_text(text, "xlsx", sheet.title)
                    self.rag(text, f"sheet:{sheet.title}")
            finally:
                book.close()
            return
        elif self.url.lower().split("?", 1)[0].endswith(".docx"):
            import zipfile
            with zipfile.ZipFile(io.BytesIO(content)) as book:
                for name in book.namelist():
                    if name.startswith("word/") and name.endswith(".xml"):
                        node = ET.fromstring(book.read(name))
                        text = "\n".join("".join(p.itertext()) for p in node.iter()
                                         if p.tag.endswith("}p"))
                        if text:
                            self.full_text(text, "docx", name)
                            self.rag(text, name)
            self.store.issue(self.document_id, self.url, "attachment_review",
                             "DOCX text retained; embedded objects and visual content need review")
            return
        elif self.url.lower().split("?", 1)[0].endswith(".zip"):
            import zipfile
            with zipfile.ZipFile(io.BytesIO(content)) as book:
                for index, entry in enumerate(book.infolist()):
                    if entry.is_dir():
                        continue
                    if entry.file_size > 1024 ** 3:
                        raise ValueError("ZIP member exceeds 1 GB; needs explicit review")
                    data = book.read(entry)
                    checksum = digest(data)
                    path = self.store.root / "objects" / checksum[:2] / checksum
                    if not path.exists():
                        atomic_bytes(path, data)
                    self.emit("attachments", f"zip:{index}", {
                        "member": entry.filename, "sha256": checksum,
                        "path": str(path.relative_to(self.store.root)), "size_bytes": len(data),
                    })
                    self.store.issue(self.document_id, self.url, "attachment_review",
                                     f"ZIP member {entry.filename} retained; extraction needs review")
            return
        self.emit("attachments", "file", {"kind": kind, "size_bytes": len(content)})
        if kind not in ("epub",):
            self.store.issue(self.document_id, self.url, "attachment_review",
                             f"{kind} retained; verify non-text content and extraction")
